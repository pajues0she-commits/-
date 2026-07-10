"""누적 이력과 연도별 판정 매트릭스를 엑셀(.xlsx)로 내보낸다. openpyxl 필요."""

from __future__ import annotations

from datetime import date

from .comparator import MatrixRow
from .storage import History


def _require_openpyxl():
    try:
        import openpyxl  # noqa: F401
        return openpyxl
    except ImportError as e:
        raise RuntimeError(
            "엑셀 내보내기에는 openpyxl이 필요합니다. 'pip install openpyxl'로 설치하세요."
        ) from e


def export_excel(
    path: str,
    history: History,
    years: list[int],
    rows: list[MatrixRow],
    department: str = "",
    author: str = "",
) -> None:
    """시트1 '누적 이력' + 시트2 '연도별 판정'을 담은 xlsx 파일을 만든다."""
    openpyxl = _require_openpyxl()
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    thin = Side(style="thin", color="999999")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor="1E3A5F")
    head_font = Font(bold=True, color="FFFFFF")
    changed_font = Font(bold=True, color="B91C1C")
    changed_fill = PatternFill("solid", fgColor="FDE8E8")
    center = Alignment(horizontal="center", vertical="center")

    wb = openpyxl.Workbook()

    # ---------------------------------------------------------- 시트1: 누적 이력
    ws = wb.active
    ws.title = "누적 이력"
    ws["A1"] = "MSDS 누적 관리 대장"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"작성부서: {department or '-'}"
    ws["A3"] = f"작성자: {author or '-'}"
    ws["A4"] = f"출력일: {date.today().isoformat()}"
    ws["A5"] = "※ 기준연도 = 도입일자의 연도 (도입일자가 없으면 MSDS 문서의 개정연도)"
    ws["A5"].font = Font(size=9, color="777777")

    headers = ["No.", "도입일자", "파일명", "기준연도", "MSDS 문서연도", "제품명", "제조사",
               "성분 수", "성분 상세 (성분명/CAS/함유량)", "작성부서", "작성자", "등록일시"]
    header_row = 6
    for col, h in enumerate(headers, start=1):
        c = ws.cell(row=header_row, column=col, value=h)
        c.font = head_font
        c.fill = head_fill
        c.border = border
        c.alignment = center

    for i, e in enumerate(
        sorted(history.entries, key=lambda x: (x.intro_date or "9999", x.added_at)),
        start=1,
    ):
        r = e.record
        detail = "; ".join(
            " / ".join(p for p in (
                ing.name or "(이름 미상)",
                f"CAS {ing.cas}" if ing.cas else "",
                f"{ing.content}%" if ing.content else "",
            ) if p)
            for ing in r.ingredients
        )
        values = [i, e.intro_date, r.filename, r.year or "?",
                  r.doc_year or "-", r.product,
                  r.manufacturer, len(r.ingredients), detail,
                  e.department, e.author, e.added_at.replace("T", " ")]
        for col, v in enumerate(values, start=1):
            c = ws.cell(row=header_row + i, column=col, value=v)
            c.border = border
            if col in (1, 2, 4, 5, 8):
                c.alignment = center

    widths = [5, 12, 26, 9, 12, 18, 22, 8, 60, 12, 10, 19]
    for col, w in enumerate(widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = w
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)

    # ---------------------------------------------------------- 시트2: 연도별 판정
    if years and rows:
        ws2 = wb.create_sheet("연도별 판정")
        ws2["A1"] = "연도별 변경 판정 (▲ = 직전 연도 대비 변경, — = 해당 연도에 없음)"
        ws2["A1"].font = Font(bold=True, size=12)
        ws2["A2"] = f"작성부서: {department or '-'}   작성자: {author or '-'}   출력일: {date.today().isoformat()}"

        header_row2 = 4
        cols2 = ["항목"] + [f"{y}년" for y in years] + ["변경 여부"]
        for col, h in enumerate(cols2, start=1):
            c = ws2.cell(row=header_row2, column=col, value=h)
            c.font = head_font
            c.fill = head_fill
            c.border = border
            c.alignment = center

        for ri, row in enumerate(rows, start=1):
            label_cell = ws2.cell(row=header_row2 + ri, column=1, value=row.label)
            label_cell.border = border
            if row.kind == "manufacturer":
                label_cell.font = Font(bold=True)
            for ci, (v, ch) in enumerate(zip(row.values, row.changed), start=2):
                c = ws2.cell(row=header_row2 + ri, column=ci,
                             value=f"▲ {v}" if ch else v)
                c.border = border
                c.alignment = center
                if ch:
                    c.font = changed_font
                    c.fill = changed_fill
            verdict = ws2.cell(row=header_row2 + ri, column=len(cols2),
                               value="변경" if row.any_changed else "변경없음")
            verdict.border = border
            verdict.alignment = center
            if row.any_changed:
                verdict.font = changed_font

        ws2.column_dimensions["A"].width = 28
        for col in range(2, len(cols2) + 1):
            ws2.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 20
        ws2.freeze_panes = ws2.cell(row=header_row2 + 1, column=2)

    wb.save(path)
