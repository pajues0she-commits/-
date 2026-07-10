"""파서·비교기 단위 테스트. GUI 없이 실행 가능: python -m unittest discover tests"""

import os
import sys
import unittest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from msds_checker.comparator import (  # noqa: E402
    CHANGED, UNCHANGED, UNKNOWN, compare_all, build_report, build_csv,
    build_matrix, build_matrix_csv,
)
from msds_checker.parser import (  # noqa: E402
    is_valid_cas, normalize_company, parse_file,
)

SAMPLES = os.path.join(os.path.dirname(__file__), "..", "samples")


def sample(name):
    return os.path.join(SAMPLES, name)


class TestCas(unittest.TestCase):
    def test_valid(self):
        for cas in ("108-88-3", "1330-20-7", "67-64-1", "64-17-5", "7732-18-5"):
            self.assertTrue(is_valid_cas(cas), cas)

    def test_invalid_checksum(self):
        self.assertFalse(is_valid_cas("108-88-4"))
        self.assertFalse(is_valid_cas("123-45-6"))


class TestNormalizeCompany(unittest.TestCase):
    def test_corp_suffixes(self):
        self.assertEqual(normalize_company("대한케미칼(주)"), normalize_company("대한케미칼 주식회사"))
        self.assertEqual(normalize_company("ABC Co., Ltd."), normalize_company("ABC"))
        self.assertNotEqual(normalize_company("대한케미칼(주)"), normalize_company("한국정밀화학 주식회사"))


class TestParser(unittest.TestCase):
    def test_parse_2022(self):
        rec = parse_file(sample("MSDS_신나A_2022.txt"))
        self.assertEqual(rec.year, 2022)  # 개정일(2022) > 작성일(2020) 우선
        self.assertEqual(rec.product, "신나 A-100")
        self.assertEqual(rec.manufacturer, "대한케미칼(주)")
        self.assertEqual([i.cas for i in rec.ingredients],
                         ["108-88-3", "1330-20-7", "67-64-1"])
        self.assertEqual(rec.ingredients[0].name, "톨루엔")
        self.assertEqual(rec.ingredients[0].content, "40 - 50")

    def test_parse_2023_added_ingredient(self):
        rec = parse_file(sample("MSDS_신나A_2023.txt"))
        self.assertEqual(rec.year, 2023)
        self.assertEqual(len(rec.ingredients), 4)

    def test_year_from_filename_fallback(self):
        import tempfile
        with tempfile.NamedTemporaryFile(
            "w", suffix="_2021.txt", delete=False, encoding="utf-8"
        ) as f:
            f.write("1. 화학제품과 회사에 관한 정보\n회사명 : 테스트사\n")
            path = f.name
        try:
            rec = parse_file(path)
            self.assertEqual(rec.year, 2021)
        finally:
            os.unlink(path)


class TestComparator(unittest.TestCase):
    def setUp(self):
        self.records = [
            parse_file(sample("MSDS_신나A_2024.txt")),
            parse_file(sample("MSDS_신나A_2022.txt")),
            parse_file(sample("MSDS_신나A_2023.txt")),
        ]
        self.comparisons, self.warnings = compare_all(self.records)

    def test_pairs_sorted_by_year(self):
        self.assertEqual([(c.year_from, c.year_to) for c in self.comparisons],
                         [(2022, 2023), (2023, 2024)])

    def test_2022_to_2023_ingredient_changed_manufacturer_same(self):
        c = self.comparisons[0]
        self.assertEqual(c.manufacturer_verdict, UNCHANGED)
        self.assertEqual(c.ingredient_verdict, CHANGED)
        self.assertEqual([i.cas for i in c.ingredient_diff.added], ["64-17-5"])
        self.assertEqual(c.ingredient_diff.removed, [])
        # 톨루엔 함유량 40-50 → 30-40
        changed = [(a.cas, a.content, b.content) for a, b in c.ingredient_diff.content_changed]
        self.assertEqual(changed, [("108-88-3", "40 - 50", "30 - 40")])

    def test_2023_to_2024_manufacturer_changed_ingredient_same(self):
        c = self.comparisons[1]
        self.assertEqual(c.manufacturer_verdict, CHANGED)
        self.assertIn("대한케미칼", c.manufacturer_detail)
        self.assertIn("한국정밀화학", c.manufacturer_detail)
        self.assertEqual(c.ingredient_verdict, UNCHANGED)

    def test_single_year_warns(self):
        comps, warns = compare_all([parse_file(sample("MSDS_신나A_2022.txt"))])
        self.assertEqual(comps, [])
        self.assertTrue(any("2개 이상" in w for w in warns))

    def test_matrix(self):
        years, rows, _ = build_matrix(self.records)
        self.assertEqual(years, [2022, 2023, 2024])
        # 첫 행은 제조사, 2024년에만 변경 표시
        mfr = rows[0]
        self.assertEqual(mfr.label, "제조사")
        self.assertEqual(mfr.values,
                         ["대한케미칼(주)", "대한케미칼(주)", "한국정밀화학 주식회사"])
        self.assertEqual(mfr.changed, [False, False, True])
        by_label = {r.label: r for r in rows}
        # 톨루엔: 2023년 함유량 변경
        tol = by_label["톨루엔 (108-88-3)"]
        self.assertEqual(tol.values, ["40 - 50", "30 - 40", "30 - 40"])
        self.assertEqual(tol.changed, [False, True, False])
        # 에탄올: 2022년 없음 → 2023년 추가
        eth = by_label["에탄올 (64-17-5)"]
        self.assertEqual(eth.values, ["—", "5 - 10", "5 - 10"])
        self.assertEqual(eth.changed, [False, True, False])
        # 크실렌: 변경 없음
        xyl = by_label["크실렌 (1330-20-7)"]
        self.assertFalse(xyl.any_changed)

    def test_matrix_csv(self):
        years, rows, _ = build_matrix(self.records)
        csv_text = build_matrix_csv(years, rows)
        lines = csv_text.strip().splitlines()
        self.assertEqual(lines[0], "항목,2022년,2023년,2024년,변경 여부")
        self.assertEqual(len(lines), 1 + len(rows))
        self.assertIn("▲ 한국정밀화학 주식회사", csv_text)

    def test_report_and_csv(self):
        report = build_report(self.records, self.comparisons, self.warnings)
        self.assertIn("2022년 → 2023년", report)
        self.assertIn("변경", report)
        csv_text = build_csv(self.comparisons)
        self.assertIn("제조사 판정", csv_text)
        self.assertEqual(len(csv_text.strip().splitlines()), 3)  # 헤더 + 2행


class TestStorage(unittest.TestCase):
    def setUp(self):
        import tempfile
        self.tmpdir = tempfile.mkdtemp()
        os.environ["MSDS_DATA_DIR"] = self.tmpdir

    def tearDown(self):
        import shutil
        os.environ.pop("MSDS_DATA_DIR", None)
        shutil.rmtree(self.tmpdir, ignore_errors=True)

    def test_add_load_remove(self):
        from msds_checker import storage
        history = storage.load()
        self.assertEqual(history.entries, [])
        rec = parse_file(sample("MSDS_신나A_2022.txt"))
        entry = storage.add_record(
            history, rec, intro_date="2022-03-15",
            department="안전보건팀", author="홍길동",
        )
        # 파일 사본이 저장소로 복사됨
        self.assertTrue(rec.path.startswith(self.tmpdir))
        self.assertTrue(os.path.exists(rec.path))
        # 다시 로드해도 유지됨 (누적 관리)
        loaded = storage.load()
        self.assertEqual(len(loaded.entries), 1)
        e = loaded.entries[0]
        self.assertEqual(e.intro_date, "2022-03-15")
        self.assertEqual(e.department, "안전보건팀")
        self.assertEqual(e.author, "홍길동")
        self.assertEqual(e.record.year, 2022)
        self.assertEqual(e.record.manufacturer, "대한케미칼(주)")
        self.assertEqual(len(e.record.ingredients), 3)
        self.assertEqual(loaded.department, "안전보건팀")
        # 제거하면 사본도 삭제됨
        storage.remove_entries(loaded, {entry.id})
        self.assertEqual(storage.load().entries, [])
        self.assertFalse(os.path.exists(rec.path))


class TestExporter(unittest.TestCase):
    def setUp(self):
        import tempfile
        self.tmpdir = tempfile.mkdtemp()
        os.environ["MSDS_DATA_DIR"] = self.tmpdir

    def tearDown(self):
        import shutil
        os.environ.pop("MSDS_DATA_DIR", None)
        shutil.rmtree(self.tmpdir, ignore_errors=True)

    def test_export_excel(self):
        try:
            import openpyxl
        except ImportError:
            self.skipTest("openpyxl 미설치")
        from msds_checker import storage
        from msds_checker.exporter import export_excel
        history = storage.load()
        for name, intro in (
            ("MSDS_신나A_2022.txt", "2022-03-15"),
            ("MSDS_신나A_2023.txt", "2023-02-01"),
            ("MSDS_신나A_2024.txt", "2024-01-20"),
        ):
            storage.add_record(
                history, parse_file(sample(name)), intro_date=intro,
                department="안전보건팀", author="홍길동", copy_file=False,
            )
        years, rows, _ = build_matrix(history.records())
        path = os.path.join(self.tmpdir, "out.xlsx")
        export_excel(path, history, years, rows,
                     department="안전보건팀", author="홍길동")

        wb = openpyxl.load_workbook(path)
        self.assertEqual(wb.sheetnames, ["누적 이력", "연도별 판정"])
        ws = wb["누적 이력"]
        self.assertEqual(ws["B7"].value, "2022-03-15")       # 첫 데이터 행 도입일자
        self.assertEqual(ws["I7"].value, "안전보건팀")
        self.assertEqual(ws["J7"].value, "홍길동")
        ws2 = wb["연도별 판정"]
        self.assertEqual(ws2["A5"].value, "제조사")
        self.assertEqual(ws2["D5"].value, "▲ 한국정밀화학 주식회사")
        self.assertEqual(ws2["E5"].value, "변경")


if __name__ == "__main__":
    unittest.main()
